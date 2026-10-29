from datetime import date, time
from decimal import Decimal as D
from pathlib import Path

import pytest

from tradein import matrix
from tradein.rules import describe_rule, load_rules, resolve, rules_of_group

TXN = date(2026, 10, 29)

# Ma trận kỳ vọng, đối chiếu từ A2, Excel mapping và PDF TB 442 (ba nguồn khớp nhau).
EXPECTED_CONVERSION = {
    1: [2, 3, 4, 5, 6, 7], 2: [2, 5], 3: [2, 3, 5], 4: [2, 3, 4, 5], 5: [2, 5],
    6: [6], 7: [6, 7], 8: [8], 9: [9], 10: [10], 11: [], 12: [], 13: [],
}


def test_conversion_matrix_matches_documents():
    assert matrix.conversion() == EXPECTED_CONVERSION


def test_all_groups_have_rules_and_schedule():
    for g in range(1, 14):
        assert rules_of_group(g), f"nhóm {g} chưa có dòng hàng"
        assert matrix.schedule_for_group(g)["rows"]


def test_schedule_percentages_sum_to_100():
    for key, sch in matrix._raw()["schedules"].items():
        assert sum(r["pct"] for r in sch["rows"]) == 100, key


def test_payment_group_schedule_assignment():
    assert matrix.groups()[1]["schedule"] == "T2"
    for g in (2, 5, 6):
        assert matrix.groups()[g]["schedule"] == "T120"
    for g in (3, 4, 7, 8, 9, 10, 11, 12, 13):
        assert matrix.groups()[g]["schedule"] == "T7"


def test_rules_ids_unique_and_valid():
    rules = load_rules()
    assert len(rules) >= 40
    for r in rules.values():
        assert 1 <= r.payment_group <= 13
        assert describe_rule(r)


def test_buyback_window():
    assert not matrix.in_buyback_window(time(14, 59))
    assert matrix.in_buyback_window(time(15, 0))
    assert matrix.in_buyback_window(time(16, 59))
    assert not matrix.in_buyback_window(time(17, 0))


def test_pearl_freshwater_changes_on_5_jan_2026():
    r = load_rules()["G2-NT-NGOT"]
    before = resolve(r, date(2026, 1, 4), TXN)
    after = resolve(r, date(2026, 1, 5), TXN)
    assert (before.thu, before.doi) == (D("0.5"), D("0.7"))
    assert after.thu == "NL_THU" and after.stone_factor == 0


def test_cz_has_customer_choice_between_jan_and_mar_2026():
    r = load_rules()["G3-CZ"]
    assert resolve(r, date(2026, 2, 1), TXN).alt is not None
    assert resolve(r, date(2026, 4, 1), TXN).alt is None
    assert resolve(r, date(2026, 4, 1), TXN).thu == D("0.7")


def test_kcr_age_tiers_and_legacy():
    r = load_rules()["G5-KCR-86-10"]
    assert resolve(r, date(2025, 6, 1), date(2026, 10, 29)).thu == D("0.90")  # đã quá 1 năm
    assert resolve(r, date(2025, 11, 1), date(2026, 10, 29)).thu == D("0.93")  # trong vòng 1 năm
    assert resolve(r, date(2025, 10, 29), date(2026, 10, 29)).thu == D("0.93")  # đúng 1 năm vẫn trong vòng 1 năm
    assert resolve(r, date(2025, 10, 28), date(2026, 10, 29)).thu == D("0.90")
    assert resolve(r, date(2018, 6, 10), TXN).thu == "CAM_KET_HD"


def test_style_after_one_year_uses_material_price():
    r = load_rules()["G6-STYLE-KC"]
    assert resolve(r, date(2026, 5, 1), TXN).doi == D("0.60")
    assert resolve(r, date(2025, 1, 1), TXN).doi == "NL_THU"


def test_kc_mau_needs_approval_after_23_jul_2023():
    r = load_rules()["G5-KC-MAU"]
    assert resolve(r, date(2023, 7, 22), TXN).thu == D("0.80")
    assert resolve(r, date(2023, 7, 23), TXN).thu == "CAN_DUYET"


@pytest.mark.skipif(not Path("/mnt/user-data/uploads").exists(), reason="không có file Excel nguồn")
def test_matrix_matches_source_excel():
    import glob
    import openpyxl

    files = glob.glob("/mnt/user-data/uploads/*.xlsx")
    if not files:
        pytest.skip("không có file Excel nguồn")
    ws = openpyxl.load_workbook(files[0]).worksheets[0]
    import re

    group_rows = {}
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row, 1).value
        m = re.match(r"\s*(\d+)", str(v)) if v is not None else None
        if m and 1 <= int(m.group(1)) <= 13 and int(m.group(1)) not in group_rows:
            group_rows[int(m.group(1))] = row
    assert set(group_rows) == set(range(1, 14))
    for g, row in group_rows.items():
        got = [c - 3 for c in range(4, 14) if ws.cell(row, c).value == "R"]  # cột D..M = nhóm 1..10
        assert got == EXPECTED_CONVERSION[g], f"nhóm {g}: Excel {got} ≠ app {EXPECTED_CONVERSION[g]}"
